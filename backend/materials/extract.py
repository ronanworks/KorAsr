"""Bounded text extraction for supported classroom document formats."""

import stat
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 200_000
MAX_PDF_PAGES = 300
MAX_PDF_PAGE_STREAM_BYTES = 5 * 1024 * 1024
MAX_PPTX_SLIDES = 200
MAX_ZIP_ENTRIES = 2_000
MAX_ZIP_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 20 * 1024 * 1024
MAX_ZIP_RATIO = 100

MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain; charset=utf-8",
    ".md": "text/markdown; charset=utf-8",
}


class MaterialError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class ExtractedMaterial:
    text: str
    status: str  # ready, needs_ocr, no_text


class _TextCollector:
    def __init__(self):
        self.parts: list[str] = []
        self.chars = 0

    def add(self, text: str, label: str = "") -> None:
        text = text.strip()
        if not text:
            return
        piece = f"[{label}]\n{text}" if label else text
        if self.chars + len(piece) + 2 > MAX_EXTRACTED_CHARS:
            raise MaterialError(f"Extracted text exceeds {MAX_EXTRACTED_CHARS} characters", 413)
        self.parts.append(piece)
        self.chars += len(piece) + 2

    def result(self, empty_status: str = "no_text") -> ExtractedMaterial:
        text = "\n\n".join(self.parts)
        return ExtractedMaterial(text=text, status="ready" if text else empty_status)


def _check_zip(path: Path, extension: str) -> None:
    """Reject path tricks and oversized OOXML parts before XML parsing."""
    required = "ppt/presentation.xml" if extension == ".pptx" else "word/document.xml"
    try:
        with ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise MaterialError("Document contains too many ZIP entries", 413)
            names = set()
            total = 0
            for info in infos:
                name = info.filename
                pure = PurePosixPath(name)
                if (not name or "\\" in name or ":" in name or pure.is_absolute()
                        or ".." in pure.parts):
                    raise MaterialError("Document contains an unsafe ZIP path")
                if name in names:
                    raise MaterialError("Document contains duplicate ZIP entries")
                mode = (info.external_attr >> 16) & 0xFFFF
                if stat.S_IFMT(mode) == stat.S_IFLNK:
                    raise MaterialError("Document contains a ZIP symlink")
                if info.flag_bits & 0x1:
                    raise MaterialError("Encrypted Office documents are not supported")
                if info.file_size > MAX_ZIP_MEMBER_BYTES:
                    raise MaterialError("Document ZIP entry is too large", 413)
                if info.file_size and info.file_size / max(info.compress_size, 1) > MAX_ZIP_RATIO:
                    raise MaterialError("Document ZIP compression ratio is unsafe", 413)
                total += info.file_size
                if total > MAX_ZIP_UNCOMPRESSED_BYTES:
                    raise MaterialError("Document expands beyond the safe size limit", 413)
                names.add(name)
            if "[Content_Types].xml" not in names or required not in names:
                raise MaterialError("File is not a valid DOCX or PPTX")
    except BadZipFile as exc:
        raise MaterialError("File is not a valid DOCX or PPTX") from exc


def _extract_pdf(path: Path) -> ExtractedMaterial:
    from pypdf import PdfReader

    with path.open("rb") as source:
        signature = source.read(5)
    if signature != b"%PDF-":
        raise MaterialError("File is not a valid PDF")
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            raise MaterialError("Encrypted PDFs are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise MaterialError(f"PDF exceeds {MAX_PDF_PAGES} pages", 413)
        collector = _TextCollector()
        for number, page in enumerate(reader.pages, 1):
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > MAX_PDF_PAGE_STREAM_BYTES:
                raise MaterialError("PDF page content stream is too large", 413)
            collector.add(page.extract_text() or "", f"Page {number}")
        return collector.result("needs_ocr")
    except MaterialError:
        raise
    except Exception as exc:
        raise MaterialError("PDF text could not be extracted") from exc


def _extract_pptx(path: Path) -> ExtractedMaterial:
    from pptx import Presentation

    def shape_lines(shape, depth: int = 0):
        if depth > 4:
            raise MaterialError("Presentation contains overly nested groups", 413)
        if shape.has_text_frame and shape.text.strip():
            yield shape.text.strip()
        if shape.has_table:
            for row in shape.table.rows:
                line = " | ".join(cell.text.strip() for cell in row.cells)
                if line.strip(" |"):
                    yield line
        for child in getattr(shape, "shapes", ()):
            yield from shape_lines(child, depth + 1)

    _check_zip(path, ".pptx")
    try:
        deck = Presentation(str(path))
        if len(deck.slides) > MAX_PPTX_SLIDES:
            raise MaterialError(f"Presentation exceeds {MAX_PPTX_SLIDES} slides", 413)
        collector = _TextCollector()
        for number, slide in enumerate(deck.slides, 1):
            lines = []
            for shape in slide.shapes:
                lines.extend(shape_lines(shape))
            collector.add("\n".join(lines), f"Slide {number}")
        return collector.result()
    except MaterialError:
        raise
    except Exception as exc:
        raise MaterialError("PPTX text could not be extracted") from exc


def _extract_docx(path: Path) -> ExtractedMaterial:
    from docx import Document

    _check_zip(path, ".docx")
    try:
        doc = Document(str(path))
        collector = _TextCollector()
        for paragraph in doc.paragraphs:
            collector.add(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                collector.add(" | ".join(cell.text.strip() for cell in row.cells))
        return collector.result()
    except MaterialError:
        raise
    except Exception as exc:
        raise MaterialError("DOCX text could not be extracted") from exc


def extract_material(path: Path, extension: str) -> ExtractedMaterial:
    """Read one already size-limited local file; never run OCR or a network API."""
    if extension not in MEDIA_TYPES:
        raise MaterialError("Supported formats: PDF, PPTX, DOCX, TXT, MD", 415)
    if path.stat().st_size > MAX_FILE_BYTES:
        raise MaterialError(f"Material exceeds {MAX_FILE_BYTES // (1024 * 1024)} MiB", 413)
    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".pptx":
        return _extract_pptx(path)
    if extension == ".docx":
        return _extract_docx(path)
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise MaterialError("Text files must use UTF-8 encoding") from exc
    if "\x00" in text:
        raise MaterialError("File contains binary data instead of text")
    collector = _TextCollector()
    collector.add(text)
    return collector.result()
